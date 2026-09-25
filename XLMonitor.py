import os
import re
import time
import shutil
import csv
import atexit
from openpyxl import load_workbook

from env_config import config
from logger import logger

WATCH_DIR = config.WATCH_DIR
OUTPUT_1F_DIR = config.OUTPUT_1F_DIR
OUTPUT_CSV_DIR = config.OUTPUT_CSV_DIR
ARCHIVE_DIR = config.ARCHIVE_DIR
POLL_INTERVAL = config.POLL_INTERVAL
MAX_ARCHIVE_FILE_AGE = config.MAX_ARCHIVE_FILE_AGE
TRIM_INTERVAL = config.TRIM_INTERVAL
TEST_MODE = config.TEST_MODE
WORKING_DIR = os.path.join(config.BASE_DIR, 'working')


def replace_commas_in_parentheses(text):
    """Replace commas found inside parentheses with spaces.

    e.g. "Dimension 2d Distance (PNT19,PNT20)" ->
         "Dimension 2d Distance (PNT19 PNT20)"
    """
    if not isinstance(text, str):
        return text

    def _replace(m: "re.Match[str]") -> str:
        return f"({m.group(1).replace(',', ' ')})"

    return re.sub(r"\(([^)]*)\)", _replace, text)


def clean_row(row):
    """Return a copy of row with commas inside parentheses removed from the
    second column."""
    row = list(row)
    if len(row) > 1:
        row[1] = replace_commas_in_parentheses(row[1])
    return row


def format_csv_filename(base_name):
    """Bracket the dash-delimited DOT machine name in a CSV base filename."""
    return re.sub(
        r"-(DOT[^-]*?\d)\s*-",
        r"-[\1] -",
        base_name,
    )


def clear_working_directory():
    """Create the local workspace and remove files from prior processing."""
    os.makedirs(WORKING_DIR, exist_ok=True)

    for name in os.listdir(WORKING_DIR):
        path = os.path.join(WORKING_DIR, name)
        if os.path.isdir(path):
            shutil.rmtree(path)
        else:
            os.remove(path)


def trim_archive():
    if not ARCHIVE_DIR:
        return

    current_time = time.time()

    for filename in os.listdir(ARCHIVE_DIR):
        file_path = os.path.join(ARCHIVE_DIR, filename)
        if os.path.isfile(file_path):
            file_age = current_time - os.path.getmtime(file_path)
            if file_age > MAX_ARCHIVE_FILE_AGE * 86400:  # Convert days to seconds
                try:
                    os.remove(file_path)
                    logger.info(f"Deleted old archive file: {filename}")
                except Exception as e:
                    logger.exception(f"Error deleting {filename}: {e}")


def convert_excel_to_csv():
    for folder in [WATCH_DIR, ARCHIVE_DIR, WORKING_DIR]:
        if folder and not os.path.exists(folder):
            try:
                os.makedirs(folder)
            except Exception as e:
                logger.error(f"Error creating folder {folder}: {e}")

    for folder in [OUTPUT_1F_DIR, OUTPUT_CSV_DIR]:
        if not folder or not os.path.isdir(folder):
            logger.error(f"Output directory not available: {folder!r}")
            return

    for filename in os.listdir(WATCH_DIR):
        if filename.lower().endswith(('.xlsx', '.xls')) and not filename.startswith('~$'):
            excel_path = os.path.join(WATCH_DIR, filename)
            time.sleep(1)
            try:
                logger.info(f"Found file: {filename}. Processing...")
                base_name = os.path.splitext(filename)[0]
                csv_base_name = format_csv_filename(base_name)
                csv_path = os.path.join(OUTPUT_CSV_DIR, f"{csv_base_name}.csv")
                working_csv_path = os.path.join(
                    WORKING_DIR,
                    f"{csv_base_name}.csv",
                )
                wb = load_workbook(excel_path, data_only=True)
                ws = wb.active
                if ws is None:
                    raise ValueError(f"Workbook has no active sheet: {excel_path!r}")

                clear_working_directory()
                with open(working_csv_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow(clean_row(row))
                if TEST_MODE:
                    logger.info(f"Test mode: CSV file created at {working_csv_path}")
                else:
                    shutil.copy2(working_csv_path, csv_path)
                    os.remove(working_csv_path)
                    shutil.copy(excel_path, os.path.join(OUTPUT_1F_DIR, filename))

                shutil.move(excel_path, os.path.join(ARCHIVE_DIR, filename))

            except Exception as e:
                logger.exception(f"Error processing {filename}: {e}")


def should_exit():
    """Return True if the KILL_ME_NOW env var requests a graceful shutdown."""
    return config.kill_me_now


def run_monitor_loop():
    last_trim_time = 0
    while True:
        if should_exit():
            logger.info("KILL_ME_NOW is set. Exiting.")
            break

        try:
            convert_excel_to_csv()
        except Exception as e:
            logger.exception(f"Unexpected error during the monitor cycle: {e}")

        if time.time() - last_trim_time >= TRIM_INTERVAL:
            try:
                trim_archive()
            except Exception as e:
                logger.exception(f"Unexpected error while trimming the archive: {e}")
            last_trim_time = time.time()

        time.sleep(POLL_INTERVAL)


def main():
    exit_reason = "normal shutdown"

    def log_exit():
        logger.info("XLMonitor exiting (%s).", exit_reason)

    atexit.register(log_exit)
    logger.info("XLMonitor started.")

    try:
        run_monitor_loop()
    except BaseException:
        exit_reason = "unhandled exception"
        logger.exception("XLMonitor terminated because of an unhandled exception.")
        raise


if __name__ == "__main__":
    main()
